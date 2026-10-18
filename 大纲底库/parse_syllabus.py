# -*- coding: utf-8 -*-
"""
大纲底库结构化解析器 v2（按列位置/表头精确映射）
把 FY26寒春.xlsx 和 FY27暑秋.xlsx 清洗成生成器可调用的统一结构 JSON。
输出: syllabus_library.json
结构: {文件:{年级:{层次:{学期:[{order,title,detail,qtype,difficulty,score,stars,starred,is_review}]}}}}
"""
import openpyxl, json, re, os, traceback

UPLOAD = "/root/uploads"
OUT = "/workspace/AI教学工作台/大纲底库/syllabus_library.json"
SEM_MAP = {"寒假":"寒假","春季":"春季","暑假":"暑假","秋季":"秋季"}
DEF_SEM = {"FY26寒春":["寒假","春季"], "FY27暑秋":["暑假","秋季"]}

def cs(c):  # cell_str
    return str(c).strip() if c is not None else ""

def count_stars(t):
    return (t or "").count("★") + (t or "").count("*")

# ---------- 列映射 ----------
def fixed_colmap(file_key, sn, max_col):
    """对表头不规范的sheet用固定列位置映射；返回None表示改用表头驱动"""
    if file_key == "FY26寒春":
        if max_col <= 5:
            return {0:"order",1:"title",2:"detail",3:"qtype",4:"score"}
        return {0:"order",1:"title",2:"detail",3:"qtype",5:"score"}  # 6列,列4空
    if sn == "初二数学精进":
        return {0:"order",1:"title",2:"detail"}
    if sn in ("初二个性化","初二数学好学"):
        return {1:"order",2:"title",3:"detail"}
    if sn in ("初三个性化","初三数学好学"):
        return {0:"order",1:"title",2:"detail"}
    if sn == "初三数学精进":
        return {1:"order",2:"title",3:"detail"}
    return None  # FY27 初一(好学/精进/志高) 用表头驱动

def parse_header(row_cells):
    """解析表头行→{col:field}"""
    cm = {}
    for col, c in enumerate(row_cells):
        t = cs(c)
        if t == "": continue
        f = None
        if any(k in t for k in ["序号","讲次","课次","次数"]): f = "order"
        elif "大标题" in t or t == "内容" or t == "课次内容": f = "title"
        elif any(k in t for k in ["小标题","具体内容","具体上课内容","主要讲解","课程目标"]): f = "detail"
        elif "题型" in t or "题目类型" in t: f = "qtype"
        elif "难度" in t or "星级" in t: f = "difficulty"
        elif "分值" in t or "分数" in t: f = "score"
        elif "内容归属" in t: f = "group"
        if f: cm[col] = f
    return cm

def extract(row_cells, col_map):
    rec = {"order":"","title":"","detail":"","qtype":"","difficulty":"","score":"","stars":0,"starred":False,"is_review":False}
    for col, f in col_map.items():
        v = cs(row_cells[col]) if col < len(row_cells) else ""
        if v == "": continue
        if f == "order": rec["order"] = v
        elif f == "title": rec["title"] = v
        elif f == "detail": rec["detail"] = v
        elif f == "qtype": rec["qtype"] = v
        elif f == "difficulty": rec["difficulty"] = v
        elif f in ("score","score2"):
            rec["score"] = (rec["score"]+" / "+v) if rec["score"] else v
        # group 忽略
    # order fallback：若col_map无order，找第一个纯数字
    if not rec["order"]:
        for col in range(min(len(row_cells),6)):
            if re.fullmatch(r"\d{1,2}", cs(row_cells[col])):
                rec["order"] = cs(row_cells[col]); break
    rec["stars"] = count_stars(rec["title"]) + count_stars(rec["detail"])
    rec["starred"] = rec["stars"] > 0
    rec["is_review"] = any(k in rec["title"]+rec["detail"] for k in ["复盘","测试","复习","期中","阶段"])
    return rec

def is_header_row(cells):
    """判断是否表头行"""
    nonempty = [cs(c) for c in cells if cs(c) != ""]
    if len(nonempty) < 2: return False
    joined = " ".join(nonempty)
    return any(k in joined for k in ["序号","讲次","课次","次数","大标题","小标题","具体内容","题型","难度","分值","分数","内容归属","主要讲解","课程目标"])

def is_seg_title(cells):
    """判断是否段落标题行"""
    nonempty = [(j,cs(c)) for j,c in enumerate(cells) if cs(c) != ""]
    if not nonempty: return False
    txt = " ".join(t for _,t in nonempty)
    total = sum(len(t) for _,t in nonempty)
    if len(nonempty) > 2 or total > 25: return False
    return any(k in txt for k in ["寒假","春季","暑假","秋季","自主学习","课程大纲","课程名称","班-"])

def find_segments(ws, file_key):
    """找段落[(start_row,sem)]"""
    segs = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if is_seg_title(row):
            txt = " ".join(cs(c) for c in row if cs(c) != "")
            sem = None
            for k,v in SEM_MAP.items():
                if k in txt: sem = v; break
            segs.append([i, sem])
    defs = DEF_SEM.get(file_key, ["寒假","春季"])
    for k in range(len(segs)):
        if segs[k][1] is None:
            segs[k][1] = defs[k % 2]
    return [(r,s) for r,s in segs]

def extract_segment(ws, start_row, end_row, col_map_fixed):
    """提取一段数据；col_map_fixed为None则在该段找表头"""
    col_map = col_map_fixed
    lessons = []
    last_title = ""
    for i in range(start_row+1, end_row):
        row = list(ws[i])
        cells = [c.value for c in row]
        nonempty = [cs(c) for c in cells if cs(c) != ""]
        if not nonempty:
            if lessons: break  # 数据已开始，遇空行结束
            continue
        # 表头行统一跳过；表头驱动时建立col_map
        if is_header_row(cells):
            if col_map_fixed is None and not col_map:
                col_map = parse_header(cells)
            continue
        # 跳过残留说明行
        joined = " ".join(nonempty)
        if any(k in joined for k in ["课程介绍","适合学生","课程目标","学习大纲","课程大纲"]) and len(nonempty) <= 4 and not any(re.fullmatch(r"\d{1,2}",cs(c)) for c in cells):
            continue
        if not col_map:
            continue
        rec = extract(cells, col_map)
        if rec["title"] or rec["detail"]:
            # 跨行讲次：title空时继承上一讲标题
            if not rec["title"] and last_title:
                rec["title"] = last_title
            if rec["title"]:
                last_title = rec["title"]
            lessons.append(rec)
    return lessons

def parse_sheet(ws, file_key, sn):
    """标准两段式sheet解析"""
    max_col = ws.max_column
    cm_fixed = fixed_colmap(file_key, sn, max_col)
    segs = find_segments(ws, file_key)
    result = {}
    for idx, (start_row, sem) in enumerate(segs):
        end_row = segs[idx+1][0] if idx+1 < len(segs) else ws.max_row + 1
        # 每段独立找表头（cm_fixed为None时）；固定映射则全sheet共用
        cm = cm_fixed  # 固定映射直接用
        lessons = extract_segment(ws, start_row, end_row, cm)
        result.setdefault(sem, lessons)
    return result

def parse_jingjin_special(ws, file_key):
    """初三数学精进：含课程介绍行，表头'学习大纲/课次/课次内容/主要讲解'"""
    cm = {1:"order",2:"title",3:"detail"}
    segs = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        nonempty = [cs(c) for c in row if cs(c) != ""]
        txt = " ".join(nonempty)
        if "课程名称" in txt and ("暑假" in txt or "秋季" in txt):
            segs.append([i, "暑假" if "暑假" in txt else "秋季"])
    result = {}
    for idx, (start_row, sem) in enumerate(segs):
        end_row = segs[idx+1][0] if idx+1 < len(segs) else ws.max_row + 1
        lessons = extract_segment(ws, start_row, end_row, cm)
        result.setdefault(sem, lessons)
    return result

def parse_zhuangzhi(ws):
    headers = [cs(c) for c in next(ws.iter_rows(values_only=True))]
    sem_titles = [re.sub(r"[（(].*?[）)]","",h).strip() for h in headers if h]
    by_sem = {t: [] for t in sem_titles}
    for row in ws.iter_rows(min_row=2, values_only=True):
        for col, t in enumerate(sem_titles):
            v = cs(row[col]) if col < len(row) else ""
            if v == "" or v == "测试": continue
            by_sem[t].append({
                "order": str(len(by_sem[t])+1), "title": v, "detail": "",
                "qtype":"", "difficulty":"", "score":"", "stars":0, "starred":False,
                "is_review": "测试" in v,
            })
    return by_sem

# ============ 主流程 ============
library = {}
files = {
    "FY26寒春": f"{UPLOAD}/1792302757691040303-FY26寒春大纲.xlsx",
    "FY27暑秋": f"{UPLOAD}/1792302757743026727-FY27暑秋大纲.xlsx",
}
for file_key, fpath in files.items():
    wb = openpyxl.load_workbook(fpath, data_only=True)
    library[file_key] = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        grade = next((g for g in ["初一","初二","初三"] if g in sn), None)
        layer = next((l for l in ["好学","精进","志高","个性化","壮志"] if l in sn), None)
        if "壮志" in sn: grade, layer = "全学段", "壮志"
        if not grade: continue
        try:
            if "壮志" in sn:
                data = parse_zhuangzhi(ws)
            elif sn == "初三数学精进":
                data = parse_jingjin_special(ws, file_key)
            else:
                data = parse_sheet(ws, file_key, sn)
            library[file_key].setdefault(grade, {})[layer] = data
        except Exception:
            print(f"  [警告] {file_key}/{sn} 解析失败:")
            traceback.print_exc()
            library[file_key].setdefault(grade, {})[layer] = {"error":"parse_failed"}

print("="*50)
print("大纲底库结构化完成，统计：")
total = 0
for fk, grades in library.items():
    print(f"\n[{fk}]")
    for g, layers in grades.items():
        for l, sems in layers.items():
            if isinstance(sems, dict) and "error" not in sems:
                for s, lessons in sems.items():
                    n = len(lessons)
                    total += n
                    print(f"  {g}/{l}/{s}: {n} 讲")
print(f"\n总计: {total} 讲")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(library, f, ensure_ascii=False, indent=2)
print(f"已输出: {OUT} ({os.path.getsize(OUT)} bytes)")
